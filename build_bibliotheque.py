# -*- coding: utf-8 -*-
"""Génère le classeur « Bibliothèque KIS » (Excel 2019, sans macros)."""
import random
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.chart import DoughnutChart, LineChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.legend import Legend
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor, TwoCellAnchor
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, Font as DFont
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "Bibliotheque_KIS.xlsx"
LOGO = "assets/logo_kis.png"
TODAY = date(2026, 10, 9)          # date de référence des données d'exemple
random.seed(42)

# ---------------------------------------------------------------- palette KIS
BLUE = "00AEEF"       # bleu ciel KIS
BLUE2 = "29ABE2"
NAVY = "0B3C5D"
GREEN = "39B54A"
YELLOW = "FFD200"
RED = "E5484D"
ORANGE = "F59E0B"
BG = "F3F6F9"
WHITE = "FFFFFF"
LINE = "E3E8EE"
MUTED = "7A8B9A"
TEXT = "1F3346"
BLUE_T = "E3F6FD"     # teintes claires
GREEN_T = "E7F6E9"
RED_T = "FDECEC"
ORANGE_T = "FFF4DF"
CALC = "F4FAFD"       # fond des colonnes calculées
FONT = "Segoe UI"

# nombre de lignes pré-formatées par table
N_LIV, N_MEM, N_EMP, N_FOU, N_MOV = 300, 300, 1000, 200, 1000
R0 = 7  # première ligne de données des tables


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def font(size=10, bold=False, color=TEXT, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


thin = lambda c=LINE: Side(style="thin", color=c)
NOB = Side(style=None)

# =============================================================== DONNÉES EXEMPLE
CAT_LIVRES = ["Roman", "Littérature africaine", "Sciences", "Mathématiques", "Histoire-Géo",
              "Langues", "Référence", "Jeunesse", "Informatique"]
CLASSES = ["1ère Primaire", "2ème Primaire", "3ème Primaire", "4ème Primaire", "5ème Primaire",
           "6ème Primaire", "7ème EB", "8ème EB", "1ère Humanités", "2ème Humanités",
           "3ème Humanités", "4ème Humanités", "Personnel"]
TYPES_MEM = ["Élève", "Enseignant", "Personnel"]
CAT_FOU = ["Papeterie", "Écriture", "Matériel didactique", "Informatique", "Entretien"]
UNITES = ["Pièce", "Paquet", "Boîte", "Rame", "Carton", "Flacon", "Rouleau"]
TYPES_MOV = ["Entrée", "Sortie"]
OUI_NON = ["Oui", "Non"]

BOOKS = [  # titre, auteur, catégorie, éditeur, année
    ("L'Enfant noir", "Camara Laye", "Littérature africaine", "Plon", 1953),
    ("Une si longue lettre", "Mariama Bâ", "Littérature africaine", "NEA", 1979),
    ("Le Vieux Nègre et la médaille", "Ferdinand Oyono", "Littérature africaine", "Julliard", 1956),
    ("Les Soleils des indépendances", "Ahmadou Kourouma", "Littérature africaine", "Seuil", 1968),
    ("Things Fall Apart", "Chinua Achebe", "Littérature africaine", "Heinemann", 1958),
    ("Ville cruelle", "Mongo Beti", "Littérature africaine", "Présence Africaine", 1954),
    ("Le Pleurer-Rire", "Henri Lopes", "Littérature africaine", "Présence Africaine", 1982),
    ("Allah n'est pas obligé", "Ahmadou Kourouma", "Littérature africaine", "Seuil", 2000),
    ("Les Bouts de bois de Dieu", "Sembène Ousmane", "Littérature africaine", "Pocket", 1960),
    ("Petit Pays", "Gaël Faye", "Littérature africaine", "Grasset", 2016),
    ("Le Petit Prince", "Antoine de Saint-Exupéry", "Jeunesse", "Gallimard", 1943),
    ("Les Misérables", "Victor Hugo", "Roman", "Folio", 1862),
    ("Germinal", "Émile Zola", "Roman", "Folio", 1885),
    ("L'Étranger", "Albert Camus", "Roman", "Folio", 1942),
    ("Candide", "Voltaire", "Roman", "Librio", 1759),
    ("Vingt mille lieues sous les mers", "Jules Verne", "Roman", "Hetzel", 1870),
    ("Le Tour du monde en 80 jours", "Jules Verne", "Roman", "Hetzel", 1872),
    ("Antigone", "Jean Anouilh", "Roman", "La Table Ronde", 1944),
    ("Le Cid", "Pierre Corneille", "Roman", "Larousse", 1637),
    ("Le Malade imaginaire", "Molière", "Roman", "Larousse", 1673),
    ("Le Lion", "Joseph Kessel", "Roman", "Folio", 1958),
    ("Bel-Ami", "Guy de Maupassant", "Roman", "Folio", 1885),
    ("Harry Potter à l'école des sorciers", "J.K. Rowling", "Jeunesse", "Gallimard Jeunesse", 1998),
    ("Le Journal d'Anne Frank", "Anne Frank", "Jeunesse", "Livre de Poche", 1947),
    ("Matilda", "Roald Dahl", "Jeunesse", "Folio Junior", 1988),
    ("Charlie et la chocolaterie", "Roald Dahl", "Jeunesse", "Folio Junior", 1964),
    ("Kirikou et la sorcière", "Michel Ocelot", "Jeunesse", "Milan", 1998),
    ("Les Contes d'Amadou Koumba", "Birago Diop", "Jeunesse", "Présence Africaine", 1947),
    ("Mathématiques 1ère Humanités", "Collectif", "Mathématiques", "Loyola", 2021),
    ("Mathématiques 2ème Humanités", "Collectif", "Mathématiques", "Loyola", 2021),
    ("Mathématiques 3ème Humanités", "Collectif", "Mathématiques", "Loyola", 2022),
    ("Algèbre et géométrie 7ème EB", "Collectif", "Mathématiques", "Médiaspaul", 2020),
    ("Calcul mental — Primaire", "Collectif", "Mathématiques", "Médiaspaul", 2019),
    ("Physique 3ème Humanités", "Collectif", "Sciences", "Loyola", 2022),
    ("Chimie générale", "Collectif", "Sciences", "De Boeck", 2018),
    ("Biologie — Le corps humain", "Collectif", "Sciences", "Nathan", 2019),
    ("Une brève histoire du temps", "Stephen Hawking", "Sciences", "Flammarion", 1988),
    ("Sciences de la Terre", "Collectif", "Sciences", "Hachette", 2017),
    ("Histoire de la RDC", "Isidore Ndaywel", "Histoire-Géo", "Le Cri", 2009),
    ("Atlas de la RD Congo", "Collectif", "Histoire-Géo", "Jeune Afrique", 2016),
    ("Géographie du Kivu", "Collectif", "Histoire-Géo", "Médiaspaul", 2015),
    ("Histoire de l'Afrique", "Joseph Ki-Zerbo", "Histoire-Géo", "Hatier", 1972),
    ("Sapiens", "Yuval Noah Harari", "Histoire-Géo", "Albin Michel", 2015),
    ("English for Today — Book 1", "Collectif", "Langues", "Longman", 2018),
    ("English Grammar in Use", "Raymond Murphy", "Langues", "Cambridge", 2019),
    ("Bescherelle — La conjugaison", "Collectif", "Langues", "Hatier", 2019),
    ("Grammaire française 6ème", "Collectif", "Langues", "Hachette", 2018),
    ("Swahili kwa Wote", "Collectif", "Langues", "Médiaspaul", 2017),
    ("Le Petit Larousse illustré", "Larousse", "Référence", "Larousse", 2023),
    ("Le Robert de poche", "Le Robert", "Référence", "Le Robert", 2022),
    ("Dictionnaire Français-Anglais", "Collectif", "Référence", "Harrap's", 2020),
    ("Encyclopédie des sciences", "Collectif", "Référence", "Larousse", 2016),
    ("Atlas mondial", "Collectif", "Référence", "Nathan", 2021),
    ("Initiation à l'informatique", "Collectif", "Informatique", "Eyrolles", 2020),
    ("Python pour les débutants", "Collectif", "Informatique", "Eyrolles", 2022),
    ("Excel 2019 — Pas à pas", "Collectif", "Informatique", "ENI", 2019),
    ("Internet et sécurité", "Collectif", "Informatique", "ENI", 2021),
    ("Le Monde de Sophie", "Jostein Gaarder", "Roman", "Points", 1991),
    ("L'Alchimiste", "Paulo Coelho", "Roman", "J'ai lu", 1988),
    ("Chaka", "Thomas Mofolo", "Littérature africaine", "Gallimard", 1925),
]
POPULAR = {0, 1, 9, 10, 22, 28, 34, 44, 58, 13, 3}

FIRST = ["Amani", "Bahati", "Neema", "Furaha", "Grâce", "Josué", "Esther", "Jonathan", "Merveille",
         "Daniel", "Sarah", "Christian", "Divine", "Patrick", "Rachel", "David", "Gloria", "Samuel",
         "Ruth", "Emmanuel", "Joëlle", "Elie", "Naomie", "Caleb", "Prisca", "Moïse", "Ornella",
         "Isaac", "Chancelle", "Benjamin", "Dorcas", "Fiston"]
LAST = ["Mugisho", "Kahindo", "Masika", "Kavira", "Kambale", "Mbusa", "Paluku", "Kasereka",
        "Muhindo", "Bisimwa", "Cirimwami", "Baraka", "Mapendo", "Zawadi", "Lukogho", "Kyakimwa",
        "Balume", "Shamavu", "Byamungu", "Mirindi", "Kalume", "Nsimire", "Wabo", "Safari",
        "Kubuya", "Ngoy", "Ilunga", "Tshibanda", "Mukendi", "Lusenge", "Mwamba", "Bahizi"]

FOURNITURES = [  # article, catégorie, unité, stock initial, seuil, prix FC, fournisseur, conso
    ("Rame de papier A4", "Papeterie", "Rame", 40, 10, 15000, "Librairie Le Phare", 34),
    ("Cahier 48 pages", "Papeterie", "Paquet", 30, 8, 12000, "Librairie Le Phare", 12),
    ("Cahier 96 pages", "Papeterie", "Paquet", 25, 8, 18000, "Librairie Le Phare", 21),
    ("Chemise cartonnée", "Papeterie", "Paquet", 15, 5, 9000, "Papeterie Kivu", 6),
    ("Classeur à levier", "Papeterie", "Pièce", 30, 10, 6500, "Papeterie Kivu", 30),
    ("Craie blanche", "Matériel didactique", "Boîte", 60, 20, 4000, "Papeterie Kivu", 45),
    ("Craie de couleur", "Matériel didactique", "Boîte", 20, 8, 6000, "Papeterie Kivu", 9),
    ("Marqueur tableau blanc", "Écriture", "Boîte", 18, 6, 22000, "Librairie Le Phare", 14),
    ("Stylo à bille bleu", "Écriture", "Boîte", 25, 6, 8000, "Librairie Le Phare", 11),
    ("Crayon HB", "Écriture", "Boîte", 20, 5, 5000, "Papeterie Kivu", 7),
    ("Gomme", "Écriture", "Boîte", 12, 4, 4500, "Papeterie Kivu", 5),
    ("Règle 30 cm", "Matériel didactique", "Pièce", 50, 15, 1500, "Papeterie Kivu", 18),
    ("Agrafes 24/6", "Papeterie", "Boîte", 20, 5, 2500, "Papeterie Kivu", 20),
    ("Colle bâton", "Papeterie", "Pièce", 40, 10, 2000, "Papeterie Kivu", 16),
    ("Cartouche d'encre HP", "Informatique", "Pièce", 8, 3, 65000, "Goma Informatique", 6),
    ("Clé USB 32 Go", "Informatique", "Pièce", 10, 3, 25000, "Goma Informatique", 4),
    ("Savon liquide", "Entretien", "Flacon", 24, 8, 7000, "Supermarché Kivu", 14),
    ("Papier hygiénique", "Entretien", "Paquet", 30, 10, 11000, "Supermarché Kivu", 26),
    ("Balai", "Entretien", "Pièce", 10, 3, 5000, "Supermarché Kivu", 3),
    ("Film plastique couverture", "Papeterie", "Rouleau", 12, 4, 8500, "Librairie Le Phare", 9),
]


def gen_data():
    books = []
    for i, (t, a, c, e, y) in enumerate(BOOKS):
        isbn = "978-2-%04d-%04d-%d" % (random.randint(1000, 9999), random.randint(1000, 9999), random.randint(0, 9))
        cop = random.choice([3, 4, 5, 6]) if i in POPULAR else random.choice([1, 2, 2, 3, 3, 4])
        shelf = "Rayon %s%d" % ("ABCDEFGHI"[CAT_LIVRES.index(c)], random.randint(1, 4))
        books.append(["LIV-%03d" % (i + 1), isbn, t, a, c, e, y, cop, shelf])

    members = []
    used = set()
    def name():
        while True:
            n = "%s %s" % (random.choice(FIRST), random.choice(LAST))
            if n not in used:
                used.add(n)
                return n
    sexe_f = {"Neema", "Furaha", "Grâce", "Esther", "Merveille", "Sarah", "Divine", "Rachel", "Gloria",
              "Ruth", "Joëlle", "Naomie", "Prisca", "Ornella", "Chancelle", "Dorcas"}
    for i in range(40):
        n = name()
        if i < 32:
            typ, cls = "Élève", CLASSES[i % 12]
        elif i < 38:
            typ, cls = "Enseignant", "Personnel"
            n = ("M. " if n.split()[0] not in sexe_f else "Mme ") + n
        else:
            typ, cls = "Personnel", "Personnel"
        sx = "F" if n.replace("Mme ", "").replace("M. ", "").split()[0] in sexe_f else "M"
        tel = "+243 9%d %03d %04d" % (random.randint(70, 99), random.randint(0, 999), random.randint(0, 9999))
        insc = date(2025, 9, random.randint(2, 20)) if i < 34 else date(2026, 9 if i < 37 else 10, random.randint(1, 8))
        members.append(["MEM-%03d" % (i + 1), n, typ, cls, sx, tel, insc])

    weights = [5 if i in POPULAR else 1 for i in range(len(books))]
    loans = []
    open_count = {}

    def pick_book(open_loan):
        while True:
            b = random.choices(range(len(books)), weights)[0]
            if not open_loan or open_count.get(b, 0) < books[b][7] - 1:
                if open_loan:
                    open_count[b] = open_count.get(b, 0) + 1
                return b

    start = date(2025, 10, 13)
    # emprunts terminés sur l'année (moins pendant les vacances de juillet-août)
    d = start
    while d < TODAY - timedelta(days=12):
        n = 0 if d.weekday() >= 5 else (0 if d.month in (7, 8) and random.random() < .85 else random.choice([0, 0, 0, 1, 1, 2]))
        for _ in range(n):
            b = pick_book(False)
            m = random.randrange(40)
            dur = random.choice([4, 6, 7, 9, 10, 12, 13, 14, 14, 16, 19, 23])
            ret = d + timedelta(days=dur)
            if ret > TODAY:
                ret = TODAY
            loans.append([d, m, b, ret, None])
        d += timedelta(days=1)
    # 8 emprunts en retard, 10 en cours (dont 2 aujourd'hui)
    late_members = random.sample(range(32), 8)
    for k, back in enumerate([16, 18, 21, 23, 26, 29, 33, 38]):
        loans.append([TODAY - timedelta(days=back), late_members[k], pick_book(True), None, None])
    for back in [0, 0, 1, 2, 3, 6, 7, 8, 10, 13]:
        loans.append([TODAY - timedelta(days=back), random.randrange(40), pick_book(True), None, None])
    # 2 retours aujourd'hui
    for back in [9, 11]:
        loans.append([TODAY - timedelta(days=back), random.randrange(40), pick_book(False), TODAY, None])
    loans.sort(key=lambda r: r[0])
    for l in loans:
        if l[3] is not None and (l[3] - l[0]).days > 14:
            l[4] = "Oui" if random.random() < .7 else "Non"
        elif l[3] is not None:
            l[4] = None
    rows = [[l[0], members[l[1]][0], books[l[2]][0], l[3], l[4]] for l in loans]

    fourn = []
    moves = []
    for i, (a, c, u, s0, seuil, prix, f, conso) in enumerate(FOURNITURES):
        code = "FOU-%03d" % (i + 1)
        fourn.append([code, a, c, u, s0, seuil, prix, f])
        # sorties étalées sur 3 mois, éventuellement un réapprovisionnement
        remaining = conso
        while remaining > 0:
            q = min(remaining, random.randint(2, 8))
            remaining -= q
            dd = TODAY - timedelta(days=random.randint(0, 95))
            moves.append([dd, code, "Sortie", q, random.choice(["Secrétariat", "Classe de 7ème EB", "Salle des profs",
                                                              "Classe de 3ème Hum", "Direction", "Primaire", "Bibliothèque"]), ""])
        if i in (1, 3, 8, 11, 13, 16):
            moves.append([TODAY - timedelta(days=random.randint(20, 60)), code, "Entrée",
                          random.choice([10, 12, 15, 20]), f, "Réapprovisionnement"])
    moves.sort(key=lambda r: r[0])
    return books, members, rows, fourn, moves


BOOKS_D, MEMBERS_D, LOANS_D, FOURN_D, MOVES_D = gen_data()

# =============================================================== CLASSEUR
wb = Workbook()
ws_dash = wb.active
ws_dash.title = "Tableau de bord"
ws_liv = wb.create_sheet("Livres")
ws_mem = wb.create_sheet("Membres")
ws_emp = wb.create_sheet("Emprunts")
ws_fou = wb.create_sheet("Fournitures")
ws_mov = wb.create_sheet("Mouvements")
ws_rec = wb.create_sheet("Recherche")
ws_par = wb.create_sheet("Paramètres")
ws_cal = wb.create_sheet("Calculs")

MENU = [  # libellé, feuille
    ("▦   Tableau de bord", "Tableau de bord"),
    ("📚   Livres", "Livres"),
    ("👥   Membres", "Membres"),
    ("🔄   Emprunts", "Emprunts"),
    ("📦   Fournitures", "Fournitures"),
    ("↕   Mouvements de stock", "Mouvements"),
    ("🔍   Recherche", "Recherche"),
    ("⚙   Paramètres", "Paramètres"),
]


def name(n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


def rng(sheet, col, n):
    return "'%s'!$%s$%d:$%s$%d" % (sheet, col, R0, col, R0 + n - 1)


# ---- noms définis (lisibles dans les formules)
for n, s, c, cnt in [
    ("LIV_Code", "Livres", "E", N_LIV), ("LIV_Titre", "Livres", "G", N_LIV), ("LIV_Auteur", "Livres", "H", N_LIV),
    ("LIV_Cat", "Livres", "I", N_LIV), ("LIV_Ex", "Livres", "L", N_LIV), ("LIV_Emp", "Livres", "N", N_LIV),
    ("LIV_Dispo", "Livres", "O", N_LIV), ("LIV_Total", "Livres", "P", N_LIV), ("LIV_Statut", "Livres", "Q", N_LIV),
    ("LIV_Score", "Livres", "S", N_LIV), ("LIV_Rech", "Livres", "T", N_LIV),
    ("MEM_ID", "Membres", "E", N_MEM), ("MEM_Nom", "Membres", "F", N_MEM), ("MEM_Type", "Membres", "G", N_MEM),
    ("MEM_Classe", "Membres", "H", N_MEM), ("MEM_Insc", "Membres", "K", N_MEM), ("MEM_Cours", "Membres", "L", N_MEM),
    ("MEM_Statut", "Membres", "O", N_MEM), ("MEM_Rech", "Membres", "Q", N_MEM),
    ("EMP_Date", "Emprunts", "F", N_EMP), ("EMP_Membre", "Emprunts", "G", N_EMP), ("EMP_Nom", "Emprunts", "H", N_EMP),
    ("EMP_Livre", "Emprunts", "I", N_EMP), ("EMP_Titre", "Emprunts", "J", N_EMP), ("EMP_Prevu", "Emprunts", "K", N_EMP),
    ("EMP_Retour", "Emprunts", "L", N_EMP), ("EMP_Statut", "Emprunts", "M", N_EMP), ("EMP_Jours", "Emprunts", "N", N_EMP),
    ("EMP_Amende", "Emprunts", "O", N_EMP), ("EMP_Paye", "Emprunts", "P", N_EMP),
    ("EMP_ScoreRetard", "Emprunts", "R", N_EMP), ("EMP_ScoreRecent", "Emprunts", "S", N_EMP),
    ("FOU_Code", "Fournitures", "E", N_FOU), ("FOU_Article", "Fournitures", "F", N_FOU), ("FOU_Unite", "Fournitures", "H", N_FOU),
    ("FOU_Seuil", "Fournitures", "J", N_FOU), ("FOU_Stock", "Fournitures", "O", N_FOU), ("FOU_Valeur", "Fournitures", "P", N_FOU),
    ("FOU_Statut", "Fournitures", "Q", N_FOU), ("FOU_Score", "Fournitures", "S", N_FOU),
    ("MOV_Date", "Mouvements", "F", N_MOV), ("MOV_Code", "Mouvements", "G", N_MOV),
    ("MOV_Type", "Mouvements", "I", N_MOV), ("MOV_Qte", "Mouvements", "J", N_MOV),
]:
    name(n, rng(s, c, cnt))

PARAMS = [  # libellé, valeur, nom, format, aide
    ("Nom de l'établissement", "Kivu International School", "P_Ecole", "@", ""),
    ("Année scolaire", "2026-2027", "P_Annee", "@", ""),
    ("Durée d'un emprunt (jours)", 14, "P_Duree", "0", "Utilisée pour la date de retour prévue"),
    ("Amende par jour de retard (FC)", 500, "P_Amende", '#,##0 "FC"', "Montant en francs congolais"),
    ("Emprunts simultanés max. par membre", 3, "P_Max", "0", "Au-delà : statut « Limite atteinte »"),
]
for i, p in enumerate(PARAMS):
    name(p[2], "'Paramètres'!$F$%d" % (R0 + i))
LISTS = [("Catégories de livres", CAT_LIVRES, "L_CatLivres"), ("Classes", CLASSES, "L_Classes"),
         ("Types de membre", TYPES_MEM, "L_TypesMembre"), ("Catégories fournitures", CAT_FOU, "L_CatFourn"),
         ("Unités", UNITES, "L_Unites"), ("Types de mouvement", TYPES_MOV, "L_TypesMvt"), ("Oui / Non", OUI_NON, "L_OuiNon")]
LIST_LEN = 20
for j, (_, vals, n) in enumerate(LISTS):
    c = CL(9 + j)
    name(n, "'Paramètres'!$%s$%d:$%s$%d" % (c, R0, c, R0 + (len(vals) - 1 if n in ("L_TypesMvt", "L_OuiNon", "L_TypesMembre") else LIST_LEN - 1)))


# =============================================================== HELPERS
def setw(ws, widths):
    for k, w in widths.items():
        ws.column_dimensions[k].width = w


def box(ws, r1, c1, r2, c2, bg=WHITE, border=LINE):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(r, c)
            cell.fill = fill(bg)
            if border:
                cell.border = Border(left=thin(border) if c == c1 else NOB, right=thin(border) if c == c2 else NOB,
                                     top=thin(border) if r == r1 else NOB, bottom=thin(border) if r == r2 else NOB)


def put(ws, ref, value, f=None, al=None, fl=None, fmt=None, merge=None):
    cell = ws[ref]
    if merge:
        ws.merge_cells(merge)
    cell.value = value
    if f:
        cell.font = f
    if al:
        cell.alignment = al
    if fl:
        cell.fill = fl
    if fmt:
        cell.number_format = fmt
    return cell


L_AL = lambda indent=1, v="center", wrap=False: Alignment(horizontal="left", vertical=v, indent=indent, wrap_text=wrap)
C_AL = Alignment(horizontal="center", vertical="center")
R_AL = lambda indent=1: Alignment(horizontal="right", vertical="center", indent=indent)


def link(sheet, cell="A1"):
    return "#'%s'!%s" % (sheet, cell)


def add_logo(ws, row, height_px, col_px=None):
    img = XLImage(LOGO)
    h = height_px
    w = int(h * 420 / 498)
    side_px = 12 + 166 + 89  # largeur approx. colonnes A:C
    off = max(0, (side_px - w) // 2)
    img.anchor = OneCellAnchor(_from=AnchorMarker(col=0, colOff=off * 9525, row=row - 1, rowOff=6 * 9525),
                               ext=XDRPositiveSize2D(w * 9525, h * 9525))
    ws.add_image(img)


def sidebar(ws, active, last_row, menu_rows, logo_row=1, logo_h=110):
    setw(ws, {"A": 1.0, "B": 23, "C": 12, "D": 3})
    for r in range(1, last_row + 1):
        for c in (1, 2, 3):
            ws.cell(r, c).fill = fill(WHITE)
        ws.cell(r, 3).border = Border(right=thin("DCE3EA"))
    add_logo(ws, logo_row, logo_h)
    for (label, sheet), r in zip(MENU, menu_rows):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        cell = ws.cell(r, 2)
        cell.value = '=HYPERLINK("%s","%s")' % (link(sheet, "A1"), label)
        if sheet == active:
            ws.cell(r, 1).fill = fill(BLUE)
            for c in (2, 3):
                ws.cell(r, c).fill = fill(BLUE_T)
            ws.cell(r, 3).border = Border(right=thin("DCE3EA"))
            cell.font = font(10.5, True, "0089BF")
        else:
            cell.font = font(10.5, False, "3C5A73")
        cell.alignment = L_AL(1)


# =============================================================== TABLES DE DONNÉES
def data_sheet(ws, title, subtitle, cols, n_rows, data, calc, helpers, chips, add_label, add_col, tab):
    """cols: liste (en-tête, largeur, format, est_calculée). calc: {col_index: formule avec {r}}."""
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab
    last = R0 + n_rows - 1
    menu_rows = list(range(8, 8 + len(MENU)))
    sidebar(ws, ws.title, last + 2, menu_rows, logo_row=1, logo_h=112)
    ws.row_dimensions[1].height = 8
    ws.row_dimensions[2].height = 32
    ws.row_dimensions[3].height = 18
    ws.row_dimensions[4].height = 26
    ws.row_dimensions[5].height = 10
    ws.row_dimensions[6].height = 30
    ncol = len(cols)
    first_c, last_c = 5, 5 + ncol - 1
    for i, (h, w, fmt, is_calc) in enumerate(cols):
        ws.column_dimensions[CL(first_c + i)].width = w
    put(ws, "E2", title, font(20, True, NAVY), L_AL(0))
    put(ws, "E3", subtitle, font(9.5, False, MUTED), L_AL(0))
    # boutons en haut à droite (2 colonnes chacun)
    ws.merge_cells(start_row=2, start_column=last_c - 1, end_row=2, end_column=last_c)
    ws.merge_cells(start_row=2, start_column=last_c - 3, end_row=2, end_column=last_c - 2)
    b = ws.cell(2, last_c - 1)
    b.value = '=HYPERLINK("#\'%s\'!%s"&(%d+COUNTIF(%s,"?*")),"%s")' % (ws.title, add_col, R0, add_label[1], add_label[0])
    b2 = ws.cell(2, last_c - 3)
    b2.value = '=HYPERLINK("%s","⌂  Tableau de bord")' % link("Tableau de bord")
    for cell, fg, bg in ((b, WHITE, BLUE), (b2, NAVY, BLUE_T)):
        cell.font = font(10, True, fg)
        cell.alignment = C_AL
    for c in range(last_c - 3, last_c + 1):
        ws.cell(2, c).fill = fill(BLUE if c >= last_c - 1 else BLUE_T)
        ws.cell(2, c).border = Border(left=Side(style="medium", color=WHITE) if c in (last_c - 1, last_c - 3) else NOB,
                                      right=Side(style="medium", color=WHITE) if c in (last_c, last_c - 2) else NOB)
    # puces d'information
    c = first_c
    for formula, color, tint in chips:
        ws.merge_cells(start_row=4, start_column=c, end_row=4, end_column=c + 1)
        ws.cell(4, c + 1).fill = fill(tint)
        ws.cell(4, c + 1).border = Border(right=Side(style="medium", color=WHITE))
        cell = ws.cell(4, c)
        cell.value = formula
        cell.font = font(9.5, True, color)
        cell.fill = fill(tint)
        cell.alignment = C_AL
        c += 2
    # en-têtes
    for i, (h, w, fmt, is_calc) in enumerate(cols):
        cell = ws.cell(6, first_c + i, h)
        cell.font = font(9.5, True, WHITE)
        cell.fill = fill(NAVY if not is_calc else "1B5E86")
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
        cell.border = Border(right=thin("2C6A90"))
    # lignes
    for r in range(R0, last + 1):
        ws.row_dimensions[r].height = 21
        for i, (h, w, fmt, is_calc) in enumerate(cols):
            cell = ws.cell(r, first_c + i)
            cell.font = font(10, False, TEXT)
            cell.number_format = fmt
            cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            cell.border = Border(bottom=thin("EDF1F5"))
            if is_calc:
                cell.fill = fill(CALC)
                cell.value = calc[i].format(r=r)
        for hc, (formula, fmt) in helpers.items():
            cell = ws.cell(r, hc)
            cell.value = formula.format(r=r)
            cell.number_format = fmt
    for k, row in enumerate(data):
        r = R0 + k
        j = 0
        for i, (h, w, fmt, is_calc) in enumerate(cols):
            if is_calc:
                continue
            if j < len(row):
                ws.cell(r, first_c + i).value = row[j]
            j += 1
    for hc in helpers:
        ws.column_dimensions[CL(hc)].hidden = True
    ws.auto_filter.ref = "%s6:%s%d" % (CL(first_c), CL(last_c), last)
    ws.freeze_panes = "E7"
    # zébrage
    rngs = "%s%d:%s%d" % (CL(first_c), R0, CL(last_c), last)
    ZEBRA.append((ws, rngs))
    return first_c, last_c, last


ZEBRA = []  # zébrage ajouté en dernier (priorité la plus basse)
STATUS_STYLES = [  # texte -> (fond, police)
    (["Disponible", "Rendu", "En stock", "Actif", "Entrée", "Oui"], GREEN_T, "1E7B34"),
    (["En retard", "Rupture", "Indisponible", "Non"], RED_T, "C62828"),
    (["Stock faible", "Dernier exemplaire", "Rendu en retard", "Limite atteinte", "Sortie"], ORANGE_T, "A15C00"),
    (["En cours"], BLUE_T, "0077A8"),
]


def status_cf(ws, ref, top_left):
    for words, bg, fc in STATUS_STYLES:
        cond = "OR(%s)" % ",".join('%s="%s"' % (top_left, w) for w in words)
        ws.conditional_formatting.add(ref, FormulaRule(formula=[cond], fill=fill(bg),
                                                       font=Font(name=FONT, bold=True, color=fc), stopIfTrue=True))


def dv_list(ws, ref, source, msg):
    dv = DataValidation(type="list", formula1="=%s" % source, allow_blank=True, showErrorMessage=True,
                        errorTitle="Valeur non valide", error=msg)
    ws.add_data_validation(dv)
    dv.add(ref)


def dv_date(ws, ref):
    dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True, showErrorMessage=True,
                        errorTitle="Date non valide", error="Saisissez une date au format jj/mm/aaaa.")
    ws.add_data_validation(dv)
    dv.add(ref)


def dv_int(ws, ref, minv=0):
    dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1=str(minv), allow_blank=True,
                        showErrorMessage=True, errorTitle="Nombre non valide", error="Saisissez un nombre entier ≥ %d." % minv)
    ws.add_data_validation(dv)
    dv.add(ref)


DATE = "dd/mm/yyyy"
FC = '#,##0 "FC";-#,##0 "FC";"–"'

# ---------------- Livres
cols = [("Code", 11, "@", False), ("ISBN", 20, "@", False), ("Titre", 36, "@", False), ("Auteur", 24, "@", False),
        ("Catégorie", 19, "@", False), ("Éditeur", 18, "@", False), ("Année", 8, "0", False),
        ("Exemplaires", 12, "0", False), ("Emplacement", 13, "@", False), ("Empruntés", 11, "0", True),
        ("Disponibles", 12, "0", True), ("Total emprunts", 13, "0", True), ("Statut", 19, "@", True)]
calc = {9: '=IF($E{r}="","",COUNTIFS(EMP_Livre,$E{r},EMP_Retour,""))',
        10: '=IF($E{r}="","",$L{r}-$N{r})',
        11: '=IF($E{r}="","",COUNTIF(EMP_Livre,$E{r}))',
        12: '=IF($E{r}="","",IF($O{r}<=0,"Indisponible",IF(AND($O{r}=1,$L{r}>1),"Dernier exemplaire","Disponible")))'}
fc, lc, last = data_sheet(
    ws_liv, "📚  Catalogue des livres", "Saisissez dans les colonnes blanches — les colonnes bleutées se calculent automatiquement.",
    cols, N_LIV, BOOKS_D, calc, {}, [
        ('=COUNTIF(LIV_Code,"?*")&" titres"', NAVY, BLUE_T),
        ('=SUM(LIV_Ex)&" exemplaires"', NAVY, BLUE_T),
        ('=SUM(LIV_Dispo)&" disponibles"', "1E7B34", GREEN_T),
        ('=SUM(LIV_Emp)&" empruntés"', "0077A8", BLUE_T),
    ], ("＋  Nouveau livre", "LIV_Code"), "E", "00AEEF")
for r in range(R0, last + 1):  # colonnes masquées (top & recherche)
    ws_liv.cell(r, 19).value = '=IF($E%d="","",$P%d+ROW()/100000)' % (r, r)
    ws_liv.cell(r, 20).value = ('=IF(OR(Q_Rech="",$E{r}=""),"",IF(ISNUMBER(SEARCH(Q_Rech,$E{r}&" "&$F{r}&" "&$G{r}&" "&$H{r}&" "&$I{r})),'
                                'COUNT($T$6:$T{p})+1,""))').format(r=r, p=r - 1)
ws_liv.column_dimensions["S"].hidden = True
ws_liv.column_dimensions["T"].hidden = True
ws_liv.column_dimensions["R"].width = 3
status_cf(ws_liv, "Q%d:Q%d" % (R0, last), "$Q%d" % R0)
dv_list(ws_liv, "I%d:I%d" % (R0, last), "L_CatLivres", "Choisissez une catégorie de la liste (modifiable dans Paramètres).")
dv_int(ws_liv, "L%d:L%d" % (R0, last), 0)
for col in "NOP":
    for r in range(R0, last + 1):
        ws_liv["%s%d" % (col, r)].alignment = C_AL
ws_liv["L6"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for r in range(R0, last + 1):
    ws_liv["K%d" % r].alignment = C_AL
    ws_liv["L%d" % r].alignment = C_AL

# ---------------- Membres
cols = [("ID membre", 12, "@", False), ("Nom complet", 28, "@", False), ("Type", 13, "@", False),
        ("Classe", 17, "@", False), ("Sexe", 7, "@", False), ("Téléphone", 18, "@", False),
        ("Date d'inscription", 15, DATE, False), ("Emprunts en cours", 13, "0", True), ("En retard", 11, "0", True),
        ("Amendes dues", 14, FC, True), ("Statut", 17, "@", True)]
calc = {7: '=IF($E{r}="","",COUNTIFS(EMP_Membre,$E{r},EMP_Retour,""))',
        8: '=IF($E{r}="","",COUNTIFS(EMP_Membre,$E{r},EMP_Statut,"En retard"))',
        9: '=IF($E{r}="","",SUMIFS(EMP_Amende,EMP_Membre,$E{r},EMP_Paye,"<>Oui"))',
        10: '=IF($E{r}="","",IF($M{r}>0,"En retard",IF($L{r}>=P_Max,"Limite atteinte","Actif")))'}
fc, lc, last = data_sheet(
    ws_mem, "👥  Membres", "Élèves, enseignants et personnel inscrits à la bibliothèque.",
    cols, N_MEM, MEMBERS_D, calc, {}, [
        ('=COUNTIF(MEM_ID,"?*")&" membres"', NAVY, BLUE_T),
        ('=COUNTIF(MEM_Type,"Élève")&" élèves"', NAVY, BLUE_T),
        ('=COUNTIF(MEM_Type,"Enseignant")&" enseignants"', NAVY, BLUE_T),
        ('=COUNTIF(MEM_Statut,"En retard")&" en retard"', "C62828", RED_T),
    ], ("＋  Nouveau membre", "MEM_ID"), "E", "39B54A")
for r in range(R0, last + 1):
    ws_mem.cell(r, 17).value = ('=IF(OR(Q_Rech="",$E{r}=""),"",IF(ISNUMBER(SEARCH(Q_Rech,$E{r}&" "&$F{r}&" "&$H{r})),'
                                'COUNT($Q$6:$Q{p})+1,""))').format(r=r, p=r - 1)
    for col in "ILM":
        ws_mem["%s%d" % (col, r)].alignment = C_AL
ws_mem.column_dimensions["Q"].hidden = True
ws_mem.column_dimensions["P"].width = 3
status_cf(ws_mem, "O%d:O%d" % (R0, last), "$O%d" % R0)
dv_list(ws_mem, "G%d:G%d" % (R0, last), "L_TypesMembre", "Choisissez : Élève, Enseignant ou Personnel.")
dv_list(ws_mem, "H%d:H%d" % (R0, last), "L_Classes", "Choisissez une classe de la liste (modifiable dans Paramètres).")
dv_list(ws_mem, "I%d:I%d" % (R0, last), '"M,F"', "Saisissez M ou F.")
dv_date(ws_mem, "K%d:K%d" % (R0, last))

# ---------------- Emprunts
cols = [("N°", 11, "@", True), ("Date d'emprunt", 14, DATE, False), ("ID membre", 12, "@", False),
        ("Nom du membre", 24, "@", True), ("Code livre", 11, "@", False), ("Titre du livre", 32, "@", True),
        ("Retour prévu", 13, DATE, True), ("Date de retour", 14, DATE, False), ("Statut", 17, "@", True),
        ("Jours de retard", 11, '0;-0;"–"', True), ("Amende", 12, FC, True), ("Amende payée", 11, "@", False)]
calc = {0: '=IF($F{r}="","","EMP-"&TEXT(ROW()-6,"0000"))',
        3: '=IF($G{r}="","",IFERROR(INDEX(MEM_Nom,MATCH($G{r},MEM_ID,0)),"⚠ ID inconnu"))',
        5: '=IF($I{r}="","",IFERROR(INDEX(LIV_Titre,MATCH($I{r},LIV_Code,0)),"⚠ Code inconnu"))',
        6: '=IF($F{r}="","",$F{r}+P_Duree)',
        8: '=IF(OR($F{r}="",$I{r}=""),"",IF($L{r}<>"",IF($L{r}>$K{r},"Rendu en retard","Rendu"),IF(TODAY()>$K{r},"En retard","En cours")))',
        9: '=IF(OR($F{r}="",$I{r}=""),"",MAX(0,IF($L{r}<>"",$L{r},TODAY())-$K{r}))',
        10: '=IF(OR($F{r}="",$I{r}=""),"",$N{r}*P_Amende)'}
emp_data = [[d0, m, b, ret, paid] for d0, m, b, ret, paid in LOANS_D]
fc, lc, last = data_sheet(
    ws_emp, "🔄  Emprunts & retours", "Nouvel emprunt : date, ID membre, code livre. Retour : saisissez simplement la date de retour.",
    cols, N_EMP, emp_data, calc, {}, [
        ('=COUNTIFS(EMP_Livre,"?*",EMP_Retour,"")&" en cours"', "0077A8", BLUE_T),
        ('=COUNTIF(EMP_Statut,"En retard")&" en retard"', "C62828", RED_T),
        ('=COUNTIF(EMP_Statut,"Rendu*")&" rendus"', "1E7B34", GREEN_T),
        ('=FIXED(SUMIFS(EMP_Amende,EMP_Paye,"<>Oui"),0)&" FC dus"', "A15C00", ORANGE_T),
    ], ("＋  Nouvel emprunt", "EMP_Date"), "F", "0B3C5D")
for r in range(R0, last + 1):
    ws_emp.cell(r, 18).value = '=IF($M{r}="En retard",$N{r}+ROW()/100000,"")'.format(r=r)
    ws_emp.cell(r, 19).value = '=IF(OR($F{r}="",$I{r}=""),"",$F{r}+ROW()/100000)'.format(r=r)
    for col in "NP":
        ws_emp["%s%d" % (col, r)].alignment = C_AL
ws_emp.column_dimensions["R"].hidden = True
ws_emp.column_dimensions["S"].hidden = True
ws_emp.column_dimensions["Q"].width = 3
status_cf(ws_emp, "M%d:M%d" % (R0, last), "$M%d" % R0)
status_cf(ws_emp, "P%d:P%d" % (R0, last), "$P%d" % R0)
ws_emp.conditional_formatting.add("N%d:N%d" % (R0, last), FormulaRule(formula=["AND(ISNUMBER($N%d),$N%d>0)" % (R0, R0)],
                                                                      font=Font(name=FONT, bold=True, color="C62828")))
dv_date(ws_emp, "F%d:F%d" % (R0, last))
dv_date(ws_emp, "L%d:L%d" % (R0, last))
dv_list(ws_emp, "G%d:G%d" % (R0, last), "MEM_ID", "Cet ID n'existe pas dans la feuille Membres.")
dv_list(ws_emp, "I%d:I%d" % (R0, last), "LIV_Code", "Ce code n'existe pas dans la feuille Livres.")
dv_list(ws_emp, "P%d:P%d" % (R0, last), "L_OuiNon", "Choisissez Oui ou Non.")

# ---------------- Fournitures
cols = [("Code", 11, "@", False), ("Article", 28, "@", False), ("Catégorie", 19, "@", False), ("Unité", 10, "@", False),
        ("Stock initial", 11, "0", False), ("Seuil minimum", 11, "0", False), ("Prix unitaire", 13, FC, False),
        ("Fournisseur", 21, "@", False), ("Entrées", 10, "0", True), ("Sorties", 10, "0", True),
        ("Stock actuel", 11, "0", True), ("Valeur du stock", 15, FC, True), ("Statut", 15, "@", True)]
calc = {8: '=IF($E{r}="","",SUMIFS(MOV_Qte,MOV_Code,$E{r},MOV_Type,"Entrée"))',
        9: '=IF($E{r}="","",SUMIFS(MOV_Qte,MOV_Code,$E{r},MOV_Type,"Sortie"))',
        10: '=IF($E{r}="","",$I{r}+$M{r}-$N{r})',
        11: '=IF($E{r}="","",$O{r}*$K{r})',
        12: '=IF($E{r}="","",IF($O{r}<=0,"Rupture",IF($O{r}<=$J{r},"Stock faible","En stock")))'}
fc, lc, last = data_sheet(
    ws_fou, "📦  Fournitures & stocks", "Le stock actuel = stock initial + entrées − sorties (voir la feuille Mouvements).",
    cols, N_FOU, FOURN_D, calc, {}, [
        ('=COUNTIF(FOU_Code,"?*")&" articles"', NAVY, BLUE_T),
        ('=COUNTIF(FOU_Statut,"En stock")&" en stock"', "1E7B34", GREEN_T),
        ('=COUNTIF(FOU_Statut,"Stock faible")&" stock faible"', "A15C00", ORANGE_T),
        ('=COUNTIF(FOU_Statut,"Rupture")&" en rupture"', "C62828", RED_T),
    ], ("＋  Nouvel article", "FOU_Code"), "E", "F59E0B")
for r in range(R0, last + 1):
    ws_fou.cell(r, 19).value = '=IF(OR($E{r}="",$Q{r}="En stock"),"",$O{r}/MAX($J{r},1)+ROW()/100000)'.format(r=r)
    for col in "IJMNO":
        ws_fou["%s%d" % (col, r)].alignment = C_AL
ws_fou.column_dimensions["S"].hidden = True
ws_fou.column_dimensions["R"].width = 3
status_cf(ws_fou, "Q%d:Q%d" % (R0, last), "$Q%d" % R0)
ws_fou.conditional_formatting.add("O%d:O%d" % (R0, last), FormulaRule(
    formula=['AND($E%d<>"",$O%d<=$J%d)' % (R0, R0, R0)], font=Font(name=FONT, bold=True, color="C62828")))
dv_list(ws_fou, "G%d:G%d" % (R0, last), "L_CatFourn", "Choisissez une catégorie de la liste.")
dv_list(ws_fou, "H%d:H%d" % (R0, last), "L_Unites", "Choisissez une unité de la liste.")
dv_int(ws_fou, "I%d:J%d" % (R0, last), 0)

# ---------------- Mouvements
cols = [("N°", 11, "@", True), ("Date", 13, DATE, False), ("Code article", 12, "@", False), ("Article", 28, "@", True),
        ("Type", 11, "@", False), ("Quantité", 10, "0", False), ("Bénéficiaire / Fournisseur", 26, "@", False),
        ("Remarque", 28, "@", False)]
calc = {0: '=IF($F{r}="","","MVT-"&TEXT(ROW()-6,"0000"))',
        3: '=IF($G{r}="","",IFERROR(INDEX(FOU_Article,MATCH($G{r},FOU_Code,0)),"⚠ Code inconnu"))'}
mov_data = [[d0, c, t, q, b, rem] for d0, c, t, q, b, rem in MOVES_D]
fc, lc, last = data_sheet(
    ws_mov, "↕  Mouvements de stock", "Chaque entrée (achat) ou sortie (distribution) met à jour le stock automatiquement.",
    cols, N_MOV, mov_data, calc, {}, [
        ('=COUNTIF(MOV_Type,"Entrée")&" entrées"', "1E7B34", GREEN_T),
        ('=COUNTIF(MOV_Type,"Sortie")&" sorties"', "A15C00", ORANGE_T),
        ('=SUMIFS(MOV_Qte,MOV_Type,"Sortie",MOV_Date,">="&(TODAY()-30))&" unités sorties / 30 j"', NAVY, BLUE_T),
    ], ("＋  Nouveau mouvement", "MOV_Date"), "F", "7A8B9A")
for r in range(R0, last + 1):
    ws_mov["J%d" % r].alignment = C_AL
status_cf(ws_mov, "I%d:I%d" % (R0, last), "$I%d" % R0)
dv_date(ws_mov, "F%d:F%d" % (R0, last))
dv_list(ws_mov, "G%d:G%d" % (R0, last), "FOU_Code", "Ce code n'existe pas dans la feuille Fournitures.")
dv_list(ws_mov, "I%d:I%d" % (R0, last), "L_TypesMvt", "Choisissez Entrée ou Sortie.")
dv_int(ws_mov, "J%d:J%d" % (R0, last), 1)

# =============================================================== PARAMÈTRES
ws = ws_par
ws.sheet_view.showGridLines = False
ws.sheet_properties.tabColor = "7A8B9A"
sidebar(ws, "Paramètres", 40, list(range(8, 8 + len(MENU))), logo_row=1, logo_h=112)
for r, h in {1: 8, 2: 32, 3: 18, 4: 26, 5: 10, 6: 30}.items():
    ws.row_dimensions[r].height = h
setw(ws, {"E": 38, "F": 30, "G": 36})
put(ws, "E2", "⚙  Paramètres", font(20, True, NAVY), L_AL(0))
put(ws, "E3", "Modifiez les valeurs bleutées : tout le classeur s'adapte automatiquement.", font(9.5, False, MUTED), L_AL(0))
for c, h in (("E", "Paramètre"), ("F", "Valeur")):
    put(ws, "%s6" % c, h, font(9.5, True, WHITE), L_AL(1), fill(NAVY))
for i, (lab, val, n, fmt, aide) in enumerate(PARAMS):
    r = R0 + i
    ws.row_dimensions[r].height = 24
    put(ws, "E%d" % r, lab, font(10, True, TEXT), L_AL(1))
    c = put(ws, "F%d" % r, val, font(10.5, True, "0077A8"), L_AL(1), fill(BLUE_T), fmt)
    for col in "EF":
        ws["%s%d" % (col, r)].border = Border(bottom=thin())
    if aide:
        ws["E%d" % r].comment = None
        put(ws, "G%d" % r, aide, font(9, False, MUTED, italic=True), L_AL(1))
put(ws, "I4", "Listes déroulantes", font(13, True, NAVY), L_AL(0))
for r in range(R0, R0 + LIST_LEN):
    ws.row_dimensions[r].height = 22
# les listes sont en I:O à partir de la ligne 7 (noms définis)
for j, (title, vals, n) in enumerate(LISTS):
    c = 9 + j
    ws.column_dimensions[CL(c)].width = 22
    h = ws.cell(6, c, title)
    h.font = font(9.5, True, WHITE)
    h.fill = fill("1B5E86")
    h.alignment = L_AL(1, wrap=True)
    for k in range(LIST_LEN):
        cell = ws.cell(R0 + k, c, vals[k] if k < len(vals) else None)
        cell.font = font(10, False, TEXT)
        cell.alignment = L_AL(1)
        cell.border = Border(bottom=thin("EDF1F5"))
        if k < len(vals) or n not in ("L_TypesMvt", "L_OuiNon", "L_TypesMembre"):
            cell.fill = fill(BLUE_T if k < len(vals) else "FAFCFE")
ws.column_dimensions["H"].width = 4
put(ws, "E17", "Astuce", font(10, True, NAVY), L_AL(1), fill(BLUE_T))
put(ws, "E18", "Pour vider les données d'exemple : sélectionnez les lignes de données (à partir de la ligne 7) "
               "dans Livres, Membres, Emprunts, Fournitures et Mouvements, puis effacez uniquement les colonnes "
               "blanches (touche Suppr). Les colonnes bleutées contiennent des formules : ne les supprimez pas.",
    font(9.5, False, TEXT), L_AL(1, "top", True), fill(BLUE_T), merge="E18:F23")
for r in range(17, 24):
    for col in "EF":
        ws["%s%d" % (col, r)].fill = fill(BLUE_T)

# =============================================================== CALCULS (masquée)
ws = ws_cal
ws.sheet_state = "hidden"
ws["A1"], ws["B1"], ws["C1"] = "Mois", "Emprunts", "Retours"
for i in range(12):
    r = i + 2
    ws["A%d" % r] = "=DATE(YEAR(TODAY()),MONTH(TODAY())-%d,1)" % (11 - i)
    ws["A%d" % r].number_format = "mmm yy"
    ws["B%d" % r] = '=COUNTIFS(EMP_Date,">="&A%d,EMP_Date,"<"&EDATE(A%d,1))' % (r, r)
    ws["C%d" % r] = '=COUNTIFS(EMP_Retour,">="&A%d,EMP_Retour,"<"&EDATE(A%d,1))' % (r, r)
ws["E1"], ws["F1"] = "Catégorie", "Exemplaires"
for i in range(len(CAT_LIVRES)):
    r = i + 2
    ws["E%d" % r] = "='Paramètres'!I%d" % (R0 + i)
    ws["F%d" % r] = "=SUMIF(LIV_Cat,E%d,LIV_Ex)" % r
ws["H1"], ws["I1"] = "État", "Articles"
for i, s in enumerate(["En stock", "Stock faible", "Rupture"]):
    ws["H%d" % (i + 2)] = s
    ws["I%d" % (i + 2)] = '=COUNTIF(FOU_Statut,H%d)' % (i + 2)
ws["K1"] = "Emprunts ce mois"
ws["K2"] = '=COUNTIFS(EMP_Date,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1),EMP_Date,"<="&TODAY())'
ws["L1"] = "Emprunts mois précédent (même période)"
ws["L2"] = ('=COUNTIFS(EMP_Date,">="&DATE(YEAR(TODAY()),MONTH(TODAY())-1,1),'
            'EMP_Date,"<="&MIN(EDATE(TODAY(),-1),DATE(YEAR(TODAY()),MONTH(TODAY()),0)))')
ws["M1"] = "Variation"
ws["M2"] = '=IF(L2=0,0,K2/L2-1)'

# =============================================================== RECHERCHE
ws = ws_rec
ws.sheet_view.showGridLines = False
ws.sheet_properties.tabColor = "29ABE2"
sidebar(ws, "Recherche", 50, list(range(8, 8 + len(MENU))), logo_row=1, logo_h=112)
for r, h in {1: 8, 2: 32, 3: 18, 4: 12, 5: 34, 6: 14}.items():
    ws.row_dimensions[r].height = h
setw(ws, {"E": 12, "F": 34, "G": 26, "H": 20, "I": 13, "J": 19, "K": 4})
put(ws, "E2", "🔍  Recherche", font(20, True, NAVY), L_AL(0))
put(ws, "E3", "Tapez un titre, un auteur, un code, un ISBN ou un nom de membre dans la case bleue, puis appuyez sur Entrée.",
    font(9.5, False, MUTED), L_AL(0))
put(ws, "E5", "🔎", font(16, False, BLUE), C_AL, fill(BLUE_T))
box(ws, 5, 6, 5, 10, BLUE_T, BLUE)
ws.merge_cells("F5:J5")
put(ws, "F5", "Petit", font(13, True, NAVY), L_AL(1))
ws["E5"].border = Border(left=thin(BLUE), top=thin(BLUE), bottom=thin(BLUE))
ws["F5"].border = Border(top=thin(BLUE), bottom=thin(BLUE))
name("Q_Rech", "TRIM('Recherche'!$F$5)")


def result_block(ws, top, title, color, headers, n, formulas, counter):
    ws.row_dimensions[top].height = 26
    put(ws, "E%d" % top, title, font(12, True, NAVY), L_AL(0))
    put(ws, "J%d" % top, counter, font(9.5, True, color), R_AL(1))
    hr = top + 1
    ws.row_dimensions[hr].height = 24
    for i, h in enumerate(headers):
        c = ws.cell(hr, 5 + i, h)
        c.font = font(9.5, True, WHITE)
        c.fill = fill(NAVY)
        c.alignment = L_AL(1)
    for k in range(1, n + 1):
        r = hr + k
        ws.row_dimensions[r].height = 21
        for i, f in enumerate(formulas):
            c = ws.cell(r, 5 + i, f.format(k=k))
            c.font = font(10, i == 1, TEXT)
            c.alignment = L_AL(1) if i not in (4,) else C_AL
            c.border = Border(bottom=thin("EDF1F5"))
            if k % 2 == 0:
                c.fill = fill("F8FAFC")
    status_cf(ws, "J%d:J%d" % (hr + 1, hr + n), "$J%d" % (hr + 1))
    return hr + n


def lk(rng_name, rank_name, k):
    return 'IFERROR(INDEX(%s,MATCH(%d,%s,0)),"")' % (rng_name, k, rank_name)


end = result_block(ws, 8, "Livres trouvés", "0077A8", ["Code", "Titre", "Auteur", "Catégorie", "Disponibles", "Statut"], 10,
                   ["=" + lk(x, "LIV_Rech", 0).replace("0,LIV_Rech", "{k},LIV_Rech") for x in
                    ("LIV_Code", "LIV_Titre", "LIV_Auteur", "LIV_Cat", "LIV_Dispo", "LIV_Statut")],
                   '=IF(Q_Rech="","",IF(COUNT(LIV_Rech)=0,"Aucun livre trouvé",COUNT(LIV_Rech)&" résultat(s)"&IF(COUNT(LIV_Rech)>10," — 10 premiers","")))')
end = result_block(ws, end + 2, "Membres trouvés", "1E7B34", ["ID", "Nom complet", "Type", "Classe", "En cours", "Statut"], 10,
                   ["=" + lk(x, "MEM_Rech", 0).replace("0,MEM_Rech", "{k},MEM_Rech") for x in
                    ("MEM_ID", "MEM_Nom", "MEM_Type", "MEM_Classe", "MEM_Cours", "MEM_Statut")],
                   '=IF(Q_Rech="","",IF(COUNT(MEM_Rech)=0,"Aucun membre trouvé",COUNT(MEM_Rech)&" résultat(s)"&IF(COUNT(MEM_Rech)>10," — 10 premiers","")))')

# =============================================================== TABLEAU DE BORD
ws = ws_dash
ws.sheet_view.showGridLines = False
ws.sheet_view.zoomScale = 85
ws.sheet_properties.tabColor = BLUE
CARD_W, GAP_W = 5.6, 1.6
widths = {"A": 1.0, "B": 30, "C": 1.0, "D": 3}
for k in range(6):
    for j in range(4):
        widths[CL(5 + 5 * k + j)] = CARD_W
    if k < 5:
        widths[CL(9 + 5 * k)] = GAP_W
widths["AH"] = 3
setw(ws, widths)
LAST_ROW = 52
row_h = {1: 12, 2: 30, 3: 18, 4: 12, 5: 4, 6: 6, 7: 20, 8: 34, 9: 18, 10: 8, 11: 14,
         12: 26, 13: 4, 27: 8, 28: 14, 29: 14, 30: 26, 31: 20, 37: 8, 38: 14, 39: 26, 40: 20, 48: 10, 49: 14,
         50: 18, 51: 12, 52: 12}
for r in range(14, 27):
    row_h[r] = 20
for r in range(32, 37):
    row_h[r] = 22
for r in range(41, 48):
    row_h[r] = 24
for r, h in row_h.items():
    ws.row_dimensions[r].height = h
# fond gris du contenu
for r in range(1, LAST_ROW + 1):
    for c in range(4, 35):
        ws.cell(r, c).fill = fill(BG)
menu_rows = list(range(14, 22))
sidebar(ws, "Tableau de bord", LAST_ROW, menu_rows, logo_row=2, logo_h=150)
put(ws, "B12", "MENU", font(8.5, True, "9AA9B8"), L_AL(1))
# résumé du jour
box(ws, 30, 2, 37, 3, "F2FAFE", "CDEBF8")
put(ws, "B30", "☀  Résumé du jour", font(10.5, True, "0089BF"), L_AL(1), merge="B30:C30")
summary = [
    ("Emprunts aujourd'hui", '=COUNTIF(EMP_Date,TODAY())', "0"),
    ("Retours aujourd'hui", '=COUNTIF(EMP_Retour,TODAY())', "0"),
    ("À rendre aujourd'hui", '=COUNTIFS(EMP_Prevu,TODAY(),EMP_Retour,"")', "0"),
    ("Nouveaux membres (mois)", '=COUNTIFS(MEM_Insc,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1))', "0"),
    ("Livres en retard", '=COUNTIF(EMP_Statut,"En retard")', "0"),
    ("Amendes à encaisser", '=SUMIFS(EMP_Amende,EMP_Paye,"<>Oui")', '#,##0 "FC"'),
]
for i, (lab, f, fmt) in enumerate(summary):
    r = 31 + i
    alert = i >= 4
    put(ws, "B%d" % r, lab, font(9, False, "3C5A73"), L_AL(1))
    put(ws, "C%d" % r, f, font(10, True, "C62828" if alert else NAVY), R_AL(1), fmt=fmt)
put(ws, "B50", "Made by Alexandre Declaire", font(8, False, "A7B4C0", italic=True), L_AL(1))

# ---- en-tête
put(ws, "E2", "Tableau de bord", font(20, True, NAVY), Alignment(horizontal="left", vertical="center"))
put(ws, "E3", '="Bonjour 👋  Bienvenue à la bibliothèque de la "&P_Ecole&"  ·  Année scolaire "&P_Annee',
    font(9.5, False, MUTED), Alignment(horizontal="left", vertical="center"))
box(ws, 2, 25, 2, 33, WHITE, LINE)
ws.merge_cells("Y2:AG2")
put(ws, "Y2", '=TODAY()', font(10, True, NAVY), C_AL, fmt='"📅  "dddd dd mmmm yyyy')
box(ws, 2, 15, 2, 23, WHITE, LINE)
ws.merge_cells("O2:W2")
put(ws, "O2", '=HYPERLINK("%s","🔍  Rechercher un livre, un membre…")' % link("Recherche", "F5"), font(9.5, False, MUTED), L_AL(1))

# ---- cartes KPI
KPIS = [
    ("📚  Total des livres", "=SUM(LIV_Ex)", '=COUNTIF(LIV_Code,"?*")&" titres au catalogue"', BLUE, "Livres"),
    ("👥  Membres", '=COUNTIF(MEM_ID,"?*")', '=COUNTIF(MEM_Type,"Élève")&" élèves · "&COUNTIF(MEM_Type,"Enseignant")&" enseignants"', GREEN, "Membres"),
    ("🔄  Emprunts en cours", '=COUNTIFS(EMP_Livre,"?*",EMP_Retour,"")', '=TEXT(IF(SUM(LIV_Ex)=0,0,SUM(LIV_Emp)/SUM(LIV_Ex)),"0%")&" du fonds en circulation"', BLUE2, "Emprunts"),
    ("📈  Emprunts ce mois", "=Calculs!K2", '=IF(Calculs!M2>=0,"▲ ","▼ ")&TEXT(ABS(Calculs!M2),"0%")&" vs mois dernier"', NAVY, "Emprunts"),
    ("⏰  Livres en retard", '=COUNTIF(EMP_Statut,"En retard")', '=FIXED(SUMIFS(EMP_Amende,EMP_Paye,"<>Oui"),0)&" FC d\'amendes dues"', RED, "Emprunts"),
    ("📦  Alertes de stock", '=COUNTIF(FOU_Statut,"Stock faible")+COUNTIF(FOU_Statut,"Rupture")', '=COUNTIF(FOU_Statut,"Rupture")&" en rupture · "&COUNTIF(FOU_Statut,"Stock faible")&" stock faible"', ORANGE, "Fournitures"),
]
for k, (lab, val, sub, color, target) in enumerate(KPIS):
    c1, c2 = 5 + 5 * k, 8 + 5 * k
    box(ws, 5, c1, 10, c2, WHITE, LINE)
    for c in range(c1, c2 + 1):
        ws.cell(5, c).fill = fill(color)
    ws.merge_cells(start_row=7, start_column=c1, end_row=7, end_column=c2)
    ws.merge_cells(start_row=8, start_column=c1, end_row=8, end_column=c2)
    ws.merge_cells(start_row=9, start_column=c1, end_row=9, end_column=c2)
    cl = ws.cell(7, c1)
    cl.value = '=HYPERLINK("%s","%s")' % (link(target, "E6"), lab)
    cl.font = font(9.5, True, MUTED)
    cl.alignment = L_AL(1)
    cv = ws.cell(8, c1, val)
    cv.font = font(22, True, NAVY if color != RED else "C62828")
    cv.alignment = L_AL(1)
    cv.number_format = "#,##0"
    cs = ws.cell(9, c1, sub)
    cs.font = font(8.5, False, MUTED)
    cs.alignment = L_AL(1)
# tendance colorée (carte 4)
c4 = ws.cell(9, 5 + 5 * 3)
ws.conditional_formatting.add("T9", FormulaRule(formula=["Calculs!$M$2>=0"], font=Font(name=FONT, bold=True, color="1E7B34")))
ws.conditional_formatting.add("T9", FormulaRule(formula=["Calculs!$M$2<0"], font=Font(name=FONT, bold=True, color="C62828")))


def panel(r1, c1, r2, c2, title, more=None):
    box(ws, r1, c1, r2, c2, WHITE, LINE)
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r1, end_column=c2 - 3 if more else c2)
    t = ws.cell(r1, c1, title)
    t.font = font(11, True, NAVY)
    t.alignment = L_AL(1)
    if more:
        ws.merge_cells(start_row=r1, start_column=c2 - 2, end_row=r1, end_column=c2)
        m = ws.cell(r1, c2 - 2)
        m.value = '=HYPERLINK("%s","%s")' % (link(more[1], more[2]), more[0])
        m.font = font(9, True, "0089BF")
        m.alignment = R_AL(1)


def table_header(r, spans):
    for (c1, c2, label, al) in spans:
        if c2 > c1:
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        cell = ws.cell(r, c1, label)
        cell.font = font(8.5, True, MUTED)
        cell.alignment = al
        for c in range(c1, c2 + 1):
            cc = ws.cell(r, c)
            cc.border = Border(bottom=thin(LINE), left=cc.border.left, right=cc.border.right)


def table_rows(r1, n, spans, rule_cols=()):
    for k in range(n):
        r = r1 + k
        for (c1, c2, f, al, fmt, fnt) in spans:
            if c2 > c1:
                ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
            cell = ws.cell(r, c1, f.format(k=k + 1))
            cell.font = fnt
            cell.alignment = al
            cell.number_format = fmt
            if k < n - 1:
                for c in range(c1, c2 + 1):
                    cc = ws.cell(r, c)
                    cc.border = Border(bottom=thin("F0F3F6"), left=cc.border.left, right=cc.border.right)


# ---- Bande A : graphique, catégories, top 5
panel(12, 5, 27, 18, "Évolution des emprunts — 12 derniers mois", ("Voir tout →", "Emprunts", "E6"))
panel(12, 20, 27, 28, '="Livres par catégorie  ·  "&FIXED(SUM(LIV_Ex),0)&" exemplaires"')
panel(12, 30, 27, 33, "🏆 Top 5 des livres")
for k in range(5):
    r = 14 + 2 * k
    ws.merge_cells(start_row=r, start_column=30, end_row=r, end_column=32)
    ws.merge_cells(start_row=r + 1, start_column=30, end_row=r + 1, end_column=32)
    sc = "LARGE(LIV_Score,%d)" % (k + 1)
    t = ws.cell(r, 30, '=IFERROR("%d. "&INDEX(LIV_Titre,MATCH(%s,LIV_Score,0)),"")' % (k + 1, sc))
    t.font = font(9, True, TEXT)
    t.alignment = L_AL(1)
    a = ws.cell(r + 1, 30, '=IFERROR("     "&INDEX(LIV_Auteur,MATCH(%s,LIV_Score,0)),"")' % sc)
    a.font = font(8, False, MUTED)
    a.alignment = L_AL(1)
    ws.merge_cells(start_row=r, start_column=33, end_row=r + 1, end_column=33)
    n_ = ws.cell(r, 33, '=IFERROR(INDEX(LIV_Total,MATCH(%s,LIV_Score,0)),"")' % sc)
    n_.font = font(10, True, "0089BF")
    n_.alignment = Alignment(horizontal="center", vertical="center")
    n_.number_format = '0"×"'
    n_.fill = fill(BLUE_T)
    n_.border = Border(right=thin(LINE), top=Side(style="medium", color=WHITE), bottom=Side(style="medium", color=WHITE))
ws.merge_cells("AD25:AG25")
put(ws, "AD25", '=HYPERLINK("%s","Voir le catalogue →")' % link("Livres", "E6"), font(9, True, "0089BF"), L_AL(1))

# ---- Bande B : retards & emprunts récents
panel(30, 5, 37, 18, "⏰ Livres en retard", ("Voir tout →", "Emprunts", "M6"))
table_header(31, [(5, 9, "LIVRE", L_AL(1)), (10, 13, "MEMBRE", L_AL(1)), (14, 16, "PRÉVU LE", L_AL(1)), (17, 18, "RETARD", C_AL)])


def emp(col, score, k):
    return 'IFERROR(INDEX(%s,MATCH(LARGE(%s,%d),%s,0)),"")' % (col, score, k, score)


f_body = font(9.5, False, TEXT)
table_rows(32, 5, [
    (5, 9, "=" + emp("EMP_Titre", "EMP_ScoreRetard", 0).replace(",0),", ",{k}),"), L_AL(1), "@", font(9.5, True, TEXT)),
    (10, 13, "=" + emp("EMP_Nom", "EMP_ScoreRetard", 0).replace(",0),", ",{k}),"), L_AL(1), "@", f_body),
    (14, 16, "=" + emp("EMP_Prevu", "EMP_ScoreRetard", 0).replace(",0),", ",{k}),"), L_AL(1), DATE, f_body),
    (17, 18, "=" + emp("EMP_Jours", "EMP_ScoreRetard", 0).replace(",0),", ",{k}),"), C_AL, '0" j"', font(10, True, "C62828")),
])
panel(30, 20, 37, 33, "🕒 Emprunts récents", ("Voir tout →", "Emprunts", "E6"))
table_header(31, [(20, 21, "DATE", L_AL(1)), (22, 25, "MEMBRE", L_AL(1)), (26, 30, "LIVRE", L_AL(1)), (31, 33, "STATUT", C_AL)])
table_rows(32, 5, [
    (20, 21, "=" + emp("EMP_Date", "EMP_ScoreRecent", 0).replace(",0),", ",{k}),"), L_AL(1), "dd/mm", f_body),
    (22, 25, "=" + emp("EMP_Nom", "EMP_ScoreRecent", 0).replace(",0),", ",{k}),"), L_AL(1), "@", f_body),
    (26, 30, "=" + emp("EMP_Titre", "EMP_ScoreRecent", 0).replace(",0),", ",{k}),"), L_AL(1), "@", font(9.5, True, TEXT)),
    (31, 33, "=" + emp("EMP_Statut", "EMP_ScoreRecent", 0).replace(",0),", ",{k}),"), C_AL, "@", font(9, True, TEXT)),
])
status_cf(ws, "AE32:AE36", "$AE32")

# ---- Bande C : alertes stock, état des stocks, actions rapides
panel(39, 5, 48, 18, "📦 Fournitures en alerte", ("Voir le stock →", "Fournitures", "E6"))
table_header(40, [(5, 9, "ARTICLE", L_AL(1)), (10, 11, "UNITÉ", L_AL(1)), (12, 13, "STOCK", C_AL), (14, 16, "SEUIL MIN.", C_AL), (17, 18, "STATUT", C_AL)])


def fou(col, k):
    return 'IFERROR(INDEX(%s,MATCH(SMALL(FOU_Score,%d),FOU_Score,0)),"")' % (col, k)


table_rows(41, 6, [
    (5, 9, "=" + fou("FOU_Article", 0).replace(",0),", ",{k}),"), L_AL(1), "@", font(9.5, True, TEXT)),
    (10, 11, "=" + fou("FOU_Unite", 0).replace(",0),", ",{k}),"), L_AL(1), "@", f_body),
    (12, 13, "=" + fou("FOU_Stock", 0).replace(",0),", ",{k}),"), C_AL, "0", font(10, True, NAVY)),
    (14, 16, "=" + fou("FOU_Seuil", 0).replace(",0),", ",{k}),"), C_AL, "0", f_body),
    (17, 18, "=" + fou("FOU_Statut", 0).replace(",0),", ",{k}),"), C_AL, "@", font(9, True, TEXT)),
])
status_cf(ws, "Q41:Q46", "$Q41")
panel(39, 20, 48, 28, "État des stocks de fournitures")
panel(39, 30, 48, 33, "⚡ Actions rapides")
ACTIONS = [
    ("＋  Nouvel emprunt", '"#\'Emprunts\'!F"&(%d+COUNTIF(EMP_Date,">0"))' % R0, True),
    ("↩  Enregistrer un retour", '"#\'Emprunts\'!L"&IFERROR(MATCH("",EMP_Retour,0)+6,7)', True),
    ("＋  Nouveau livre", '"#\'Livres\'!E"&(%d+COUNTIF(LIV_Code,"?*"))' % R0, False),
    ("＋  Nouveau membre", '"#\'Membres\'!E"&(%d+COUNTIF(MEM_ID,"?*"))' % R0, False),
    ("↕  Mouvement de stock", '"#\'Mouvements\'!F"&(%d+COUNTIF(MOV_Date,">0"))' % R0, False),
    ("🔍  Rechercher", '"#\'Recherche\'!F5"', False),
]
for i, (lab, target, primary) in enumerate(ACTIONS):
    r = 41 + i
    ws.merge_cells(start_row=r, start_column=30, end_row=r, end_column=33)
    c = ws.cell(r, 30, '=HYPERLINK(%s,"%s")' % (target, lab))
    c.font = font(9.5, True, WHITE if primary else NAVY)
    c.alignment = C_AL
    for cc in range(30, 34):
        cell = ws.cell(r, cc)
        cell.fill = fill(BLUE if primary else BLUE_T)
        cell.border = Border(top=Side(style="thick", color=WHITE), bottom=Side(style="thick", color=WHITE),
                             left=thin(LINE) if cc == 30 else Side(style="thick", color=WHITE) if False else NOB,
                             right=thin(LINE) if cc == 33 else NOB)
# pied de page
put(ws, "E50", '="Bibliothèque "&P_Ecole&"  ·  Les chiffres se mettent à jour automatiquement à partir des feuilles de données."',
    font(8.5, False, "9AA9B8", italic=True), L_AL(0))


# =============================================================== GRAPHIQUES
def px_col(w):
    return int(w * 7 + 5)


def col_off(ws_, col_idx):
    return 0


def anchor(c1, r1, c2, r2, dx1=6, dy1=2, dx2=-6, dy2=-4):
    """Ancre à deux cellules (indices 1-based, coin bas/droite exclu)."""
    a = TwoCellAnchor()
    a._from = AnchorMarker(col=c1 - 1, colOff=dx1 * 9525, row=r1 - 1, rowOff=dy1 * 9525)
    a.to = AnchorMarker(col=c2 - 1, colOff=max(0, dx2) * 9525, row=r2 - 1, rowOff=max(0, dy2) * 9525)
    return a


def axis_text(size=800, color=MUTED):
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=size, solidFill=color,
                                                                                     latin=DFont(typeface=FONT))),
                                 endParaRPr=CharacterProperties(sz=size))])


def no_line():
    return GraphicalProperties(ln=LineProperties(noFill=True))


# courbe Emprunts / Retours
lc_ = LineChart()
lc_.add_data(Reference(ws_cal, min_col=2, min_row=1, max_col=3, max_row=13), titles_from_data=True)
lc_.set_categories(Reference(ws_cal, min_col=1, min_row=2, max_row=13))
for s, color in zip(lc_.series, [BLUE, GREEN]):
    s.smooth = False
    s.graphicalProperties.line.solidFill = color
    s.graphicalProperties.line.width = 31750
    s.marker.symbol = "circle"
    s.marker.size = 6
    s.marker.graphicalProperties = GraphicalProperties(solidFill=color)
    s.marker.graphicalProperties.line.solidFill = WHITE
lc_.legend.position = "t"
lc_.legend.txPr = axis_text(900, TEXT)
lc_.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E9EEF3"))
lc_.y_axis.spPr = no_line()
lc_.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill="D5DDE5"))
lc_.x_axis.number_format = "mmm"
lc_.x_axis.txPr = axis_text()
lc_.y_axis.txPr = axis_text()
lc_.x_axis.delete = False
lc_.y_axis.delete = False
lc_.y_axis.number_format = "0"
lc_.y_axis.scaling.min = 0
lc_.x_axis.majorTickMark = "none"
lc_.y_axis.majorTickMark = "none"
lc_.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
lc_.plot_area.graphicalProperties = GraphicalProperties(noFill=True)
lc_.anchor = anchor(5, 13, 19, 27, 6, 2, 0, 0)
ws.add_chart(lc_)

# anneau catégories
PALETTE = [BLUE, NAVY, GREEN, YELLOW, "7FD3F5", "2E7D9A", "8BD49A", "F28C6A", "9AA9B8", "5C6BC0"]


def donut(data_ref, cats_ref, colors, anchor_, legend_pos="r", pct=True):
    d = DoughnutChart(holeSize=62)
    d.add_data(data_ref, titles_from_data=True)
    d.set_categories(cats_ref)
    s = d.series[0]
    for i, col in enumerate(colors):
        pt = DataPoint(idx=i)
        pt.graphicalProperties = GraphicalProperties(solidFill=col)
        pt.graphicalProperties.line.solidFill = WHITE
        pt.graphicalProperties.line.width = 19050
        s.dPt.append(pt)
    from openpyxl.chart.label import DataLabelList
    s.dLbls = DataLabelList()
    s.dLbls.showPercent = pct
    s.dLbls.showVal = False
    s.dLbls.showCatName = False
    s.dLbls.showSerName = False
    s.dLbls.showLeaderLines = False
    s.dLbls.showLegendKey = False
    s.dLbls.txPr = axis_text(800, WHITE)
    d.legend.position = legend_pos
    d.legend.txPr = axis_text(850, TEXT)
    d.title = None
    d.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    d.anchor = anchor_
    return d


ws.add_chart(donut(Reference(ws_cal, min_col=6, min_row=1, max_row=1 + len(CAT_LIVRES)),
                   Reference(ws_cal, min_col=5, min_row=2, max_row=1 + len(CAT_LIVRES)),
                   PALETTE, anchor(20, 13, 29, 27, 6, 2, 0, 0)))
ws.add_chart(donut(Reference(ws_cal, min_col=9, min_row=1, max_row=4),
                   Reference(ws_cal, min_col=8, min_row=2, max_row=4),
                   [GREEN, ORANGE, RED], anchor(20, 40, 29, 48, 6, 2, 0, 0)))

# =============================================================== FINITIONS
for w_, r_ in ZEBRA:
    w_.conditional_formatting.add(r_, FormulaRule(formula=['AND($E%d<>"",MOD(ROW(),2)=0)' % R0], fill=fill("F8FAFC")))
for w in wb.worksheets:
    w.sheet_view.zoomScale = w.sheet_view.zoomScale or 100
    w.page_setup.orientation = "landscape"
    w.page_setup.fitToWidth = 1
    w.page_setup.fitToHeight = 0
    w.sheet_properties.pageSetUpPr.fitToPage = True
ws_dash.page_setup.fitToHeight = 1
for w_, key, lastcol in ((ws_liv, "LIV_Code", "Q"), (ws_mem, "MEM_ID", "O"), (ws_emp, "EMP_Date", "P"),
                         (ws_fou, "FOU_Code", "Q"), (ws_mov, "MOV_Date", "L")):
    crit = '">0"' if "Date" in key else '"?*"'
    dn = DefinedName("_xlnm.Print_Area", localSheetId=wb.worksheets.index(w_),
                     attr_text="OFFSET('%s'!$A$1,0,0,%d+COUNTIF(%s,%s),COLUMN('%s'!$%s$1))" % (w_.title, R0 - 1, key, crit, w_.title, lastcol))
    w_.defined_names["_xlnm.Print_Area"] = dn
    w_.print_title_rows = "6:6"
ws_dash.print_area = "A1:AH52"
wb.active = 0
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("OK", OUT, len(LOANS_D), "emprunts", len(MOVES_D), "mouvements")
